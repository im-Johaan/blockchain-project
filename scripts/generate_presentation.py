import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Enterprise Dark Palette
    C_BG = RGBColor(11, 19, 41)          # Navy 900 #0B1329
    C_CARD_BG = RGBColor(15, 26, 56)     # Navy 850 #0F1A38
    C_CARD_BORDER = RGBColor(30, 49, 100)# Navy 700 #1E3164
    C_PRIMARY = RGBColor(59, 130, 246)   # Blue 500 #3B82F6
    C_CYAN = RGBColor(6, 182, 212)       # Cyan 500 #06B6D4
    C_EMERALD = RGBColor(16, 185, 129)   # Emerald 500 #10B981
    C_AMBER = RGBColor(245, 158, 11)     # Amber 500 #F59E0B
    C_ROSE = RGBColor(244, 63, 94)       # Rose 500 #F43F5E
    C_WHITE = RGBColor(255, 255, 255)
    C_TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
    C_TEXT_LIGHT = RGBColor(226, 232, 240) # Slate 200

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="DATALEAK LEDGER | BLOCKCHAIN & FORENSIC STEGANOGRAPHY"):
        # Category label
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_CYAN
        p_cat.font.name = "Inter"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_WHITE
        p_title.font.name = "Inter"

    def add_card(slide, left, top, width, height, title="", subtitle="", border_color=C_CARD_BORDER, bg_color=C_CARD_BG):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.2)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)

        if title:
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = C_WHITE
            p.font.name = "Inter"

        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.size = Pt(9.5)
            p_sub.font.color.rgb = C_CYAN
            p_sub.font.name = "Inter"
            p_sub.space_after = Pt(8)

        return shape, tf

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Accent Top Banner
    banner = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.73), Inches(0.06))
    banner.fill.solid()
    banner.fill.fore_color.rgb = C_PRIMARY
    banner.line.fill.background()

    # Title box
    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.73), Inches(3.0))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "COLLEGE SEMESTER PROJECT DEFENSE"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = C_CYAN
    p0.font.name = "Inter"
    p0.space_after = Pt(10)

    p1 = tf1.add_paragraph()
    p1.text = "DataLeak Ledger"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.font.name = "Inter"

    p2 = tf1.add_paragraph()
    p2.text = "Blockchain-Backed Document Audit & Forensic Steganography Platform"
    p2.font.size = Pt(20)
    p2.font.bold = False
    p2.font.color.rgb = C_PRIMARY
    p2.font.name = "Inter"
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Solving Privileged Database Log Tampering & Untraceable Copy-Paste Exfiltration"
    p3.font.size = Pt(13)
    p3.font.color.rgb = C_TEXT_MUTED
    p3.font.name = "Inter"
    p3.space_before = Pt(8)

    # Bottom Metadata Card
    c_meta, tf_meta = add_card(s1, Inches(0.8), Inches(4.8), Inches(11.73), Inches(1.8),
                               title="Project & Architecture Specifications",
                               subtitle="EVM Smart Contracts (Solidity 0.8.24) + Zero-Width Unicode Steganography + React")
    
    p_m1 = tf_meta.add_paragraph()
    p_m1.text = "• Author / Presenter: Johaan Sam Sanil  |  Course: Blockchain Architecture & Cryptography"
    p_m1.font.size = Pt(11)
    p_m1.font.color.rgb = C_TEXT_LIGHT
    
    p_m2 = tf_meta.add_paragraph()
    p_m2.text = "• Core Stack: Solidity 0.8.24, Hardhat, ethers.js v6, Node.js CSPRNG, React 18, Tailwind CSS"
    p_m2.font.size = Pt(11)
    p_m2.font.color.rgb = C_TEXT_LIGHT
    p_m2.space_before = Pt(4)

    p_m3 = tf_meta.add_paragraph()
    p_m3.text = "• Verification Status: 26 Automated Unit Tests Passing (20 Smart Contract + 6 Forensic Tests)"
    p_m3.font.size = Pt(11)
    p_m3.font.color.rgb = C_EMERALD
    p_m3.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: THE DUAL VULNERABILITY (PROBLEM STATEMENT)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Problem Statement: The Failure of Conventional Document Systems")

    c2_1, tf2_1 = add_card(s2, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="Vulnerability 1: Privileged Log Tampering",
                           subtitle="The Centralized Database Paradox (PostgreSQL / SQL)",
                           border_color=C_ROSE)
    
    items_v1 = [
        ("Single Point of Trust: ", "Audit logs reside in traditional SQL tables."),
        ("Rogue Administrator Risk: ", "Anyone with DBA or root credentials can silently execute SQL updates or deletes without leaving cryptographic evidence."),
        ("Arbitrary Backdating: ", "Timestamps can be modified retroactively to falsify corporate compliance."),
        ("No Non-Repudiation: ", "Conventional access logs fail legal scrutiny because they lack cryptographic signatures.")
    ]
    for label, desc in items_v1:
        p = tf2_1.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    c2_2, tf2_2 = add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="Vulnerability 2: Untraceable Text Leaks",
                           subtitle="The Copy-Paste Exfiltration Gap",
                           border_color=C_AMBER)
    
    items_v2 = [
        ("Visual Watermarks Fail: ", "Diagonal stamps, headers, and barcodes disappear the moment text is copied into a clipboard."),
        ("Indistinguishable Copies: ", "Conventional servers issue identical document files to all authorized users."),
        ("Attribution Failure: ", "When a leaked paragraph surfaces in an email or public forum, standard access logs show who downloaded the file, but CANNOT prove which issued copy it originated from."),
        ("The Need: ", "Invisibly unique fingerprinting that survives copy-pasting and binds mathematically to an on-chain receipt.")
    ]
    for label, desc in items_v2:
        p = tf2_2.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    # =========================================================================
    # SLIDE 3: SYSTEM ARCHITECTURE & SOLUTION
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "System Solution: Hybrid On-Chain / Off-Chain Architecture")

    # 3 Column Cards
    col_w = Inches(3.7)
    gap = Inches(0.3)
    
    # Col 1: React Frontend
    c3_1, tf3_1 = add_card(s3, Inches(0.8), Inches(1.6), col_w, Inches(5.2),
                           title="1. React Frontend",
                           subtitle="Cybersecurity UI & Telemetry",
                           border_color=C_PRIMARY)
    c3_1_items = [
        "Vite + React 18 + TypeScript",
        "Role-aware dashboards: Admin vs Authorized Employee.",
        "On-chain access management with real-time toggle.",
        "Forensic Attribution Portal with text payload scanner.",
        "EVM Block Explorer & Chronological Audit Timeline.",
        "Instant WebSocket alerts on permission revocation."
    ]
    for item in c3_1_items:
        p = tf3_1.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_before = Pt(6)

    # Col 2: Express Backend & Forensic Engine
    c3_2, tf3_2 = add_card(s3, Inches(0.8) + col_w + gap, Inches(1.6), col_w, Inches(5.2),
                           title="2. Backend & Forensics",
                           subtitle="Atomic Enforcement Layer",
                           border_color=C_CYAN)
    c3_2_items = [
        "Express.js + TypeScript REST Gateway.",
        "Deterministic text canonicalization (CRLF/LF).",
        "Node.js CSPRNG 256-bit token generator.",
        "Zero-Width Unicode encoder (\u200B, \u200C).",
        "Ethers.js v6 contract client & Socket.IO listener.",
        "Enforces live on-chain hasAccess() before serving files."
    ]
    for item in c3_2_items:
        p = tf3_2.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_before = Pt(6)

    # Col 3: Smart Contract Core
    c3_3, tf3_3 = add_card(s3, Inches(0.8) + (col_w + gap)*2, Inches(1.6), col_w, Inches(5.2),
                           title="3. DataLeakLedger.sol",
                           subtitle="EVM Smart Contract Core",
                           border_color=C_EMERALD)
    c3_3_items = [
        "Solidity 0.8.24 + OpenZeppelin AccessControl.",
        "Authoritative single source of truth.",
        "Stores canonical SHA-256 document fingerprints.",
        "Immutable version tree (previousHash chaining).",
        "Dynamic permission mapping (hasAccess).",
        "Stores Keccak-256 issuance commitments."
    ]
    for item in c3_3_items:
        p = tf3_3.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = C_TEXT_LIGHT
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 4: WHY BLOCKCHAIN & SMART CONTRACTS?
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Why Blockchain? Decentralized Authority & Economics")

    c4_1, tf4_1 = add_card(s4, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="Blockchain as Supreme Authority",
                           subtitle="Replacing Trust in Humans with Cryptography",
                           border_color=C_PRIMARY)
    reasons = [
        ("Tamper-Proof Immutability: ", "Every block is sealed with parentHash. Modifying historical records breaks the cryptographic hash chain."),
        ("Non-Repudiation: ", "All transactions are cryptographically signed using ECDSA (secp256k1). Signers cannot deny their actions."),
        ("Deterministic State Machine: ", "EVM bytecode execution is mathematically uniform across all nodes."),
        ("Bloom Filter Auditing: ", "EVM events (FileRegistered, IssuanceRecorded) create append-only logs for sub-millisecond filtering.")
    ]
    for label, desc in reasons:
        p = tf4_1.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    c4_2, tf4_2 = add_card(s4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="Gas & Storage Economics",
                           subtitle="Why Raw Files Remain Off-Chain",
                           border_color=C_CYAN)
    econ_points = [
        ("The Gas Reality: ", "Storing 1 MB of data on Ethereum costs ~640 million gas (SSTORE opcodes). It is economically prohibitive."),
        ("The Privacy Reality: ", "All EVM contract storage is public to node operators. Storing confidential corporate text on-chain violates privacy and GDPR laws."),
        ("Our 32-Byte Solution: ", "We store only 32-byte cryptographic fingerprints (bytes32 contentHash and bytes32 commitment)."),
        ("Cost: ", "A fixed ~45,000 gas per registration regardless of document size (1 KB or 500 MB). Complete confidentiality preserved.")
    ]
    for label, desc in econ_points:
        p = tf4_2.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    # =========================================================================
    # SLIDE 5: SMART CONTRACT INTERNALS (DATALEAKLEDGER.SOL)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Smart Contract Deep-Dive: DataLeakLedger.sol")

    c5_1, tf5_1 = add_card(s5, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="Role-Based Access & Functions",
                           subtitle="OpenZeppelin AccessControl Architecture",
                           border_color=C_EMERALD)
    sc_funcs = [
        ("ADMIN_ROLE: ", "Restricted to document registrar. Manages lifecycle."),
        ("ISSUER_ROLE: ", "Restricted to trusted backend signing service for commitments."),
        ("registerDocument(): ", "Binds bytes32 documentId to canonical SHA-256 hash. Rejects duplicate IDs."),
        ("addVersion(): ", "Increments version, stores new hash, links previousHash."),
        ("grantAccess() / revokeAccess(): ", "Controls per-document permissions on-chain."),
        ("hasAccess(): ", "Constant-time O(1) live view check for downloads.")
    ]
    for label, desc in sc_funcs:
        p = tf5_1.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    c5_2, tf5_2 = add_card(s5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="Issuance Registry & Gas Optimization",
                           subtitle="Cryptographic Commitments & Solidity 0.8.24",
                           border_color=C_PRIMARY)
    sc_opt = [
        ("recordIssuance(): ", "Records issuanceCommitment, docId, version, userAddress. Crucial Guard: Reverts with AccessDenied if user has no active permission."),
        ("Commitment Uniqueness: ", "Guarantees no two downloads share the same cryptographic token."),
        ("Custom Errors vs Strings: ", "Uses 4-byte custom error selectors (e.g. error DocumentNotFound(bytes32)) saving substantial deployment and transaction gas compared to strings."),
        ("Indexed Event Logs: ", "Emits FileRegistered, VersionAdded, AccessGranted, AccessRevoked, IssuanceRecorded.")
    ]
    for label, desc in sc_opt:
        p = tf5_2.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    # =========================================================================
    # SLIDE 6: CANONICALIZATION & VERSION CHAINING
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Integrity Assurance: Deterministic Canonicalization & Version Tree")

    c6_1, tf6_1 = add_card(s6, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="Deterministic Canonical Hashing",
                           subtitle="Overcoming Cross-Platform Discrepancies",
                           border_color=C_CYAN)
    canon_points = [
        ("The Challenge: ", "Windows uses CRLF (\\r\\n), Mac/Linux use LF (\\n). Invisible trailing line spaces alter SHA-256 hashes completely (Avalanche effect)."),
        ("Our Canonical Pipeline: ", "1. Convert \\r\\n and \\r to standard \\n.\n2. Strip trailing whitespace from every line.\n3. Trim leading and trailing empty document newlines.\n4. Encode as UTF-8."),
        ("Security Guarantee: ", "Identical text yields the exact same SHA-256 hash regardless of OS. If a single character is modified, the hash mismatch is immediately caught.")
    ]
    for label, desc in canon_points:
        p = tf6_1.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    c6_2, tf6_2 = add_card(s6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="Immutable Version Tree",
                           subtitle="Historical Preservation via Hash Pointers",
                           border_color=C_PRIMARY)
    ver_points = [
        ("Version 1 (Initial): ", "contentHash = SHA256(v1), previousHash = 0x0."),
        ("Version 2 (Revision): ", "contentHash = SHA256(v2), previousHash = SHA256(v1)."),
        ("Version N (Chain): ", "Forms an immutable on-chain linked chain of document revisions."),
        ("Preservation Rule: ", "Historical versions are NEVER overwritten. Earlier issued document copies retain valid references to their exact original version.")
    ]
    for label, desc in ver_points:
        p = tf6_2.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    # =========================================================================
    # SLIDE 7: COMMITMENT SCHEME (ZERO-KNOWLEDGE ATTRIBUTION)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Security Engine: Cryptographic Token Commitment Scheme")

    c7_1, tf7_1 = add_card(s7, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="CSPRNG Token Generation",
                           subtitle="256-Bit True Entropy",
                           border_color=C_PRIMARY)
    token_points = [
        ("High Entropy: ", "Every download generates a fresh 256-bit (32-byte) random token T using Node.js crypto.randomBytes(32)."),
        ("No Math.random(): ", "Math.random() is pseudo-random and vulnerable to seed prediction. CSPRNG draws entropy directly from the OS kernel (/dev/urandom)."),
        ("Uniqueness: ", "Collision probability across 256 bits is 1 in 2^128 (virtually impossible)."),
        ("Volatile Life: ", "The raw token T is held in volatile memory only during PDF/Text rendering and never logged.")
    ]
    for label, desc in token_points:
        p = tf7_1.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    c7_2, tf7_2 = add_card(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="The One-Way Commitment (C = Keccak256(T))",
                           subtitle="Zero-Knowledge Attribution Guarantee",
                           border_color=C_EMERALD)
    commit_points = [
        ("The Mathematical Commitment: ", "C = Keccak256(T). Only C is sent to the blockchain."),
        ("Pre-Image Resistance: ", "It is computationally infeasible to deduce T from C. Even an attacker with complete access to the blockchain cannot extract T."),
        ("Framing Prevention: ", "If raw tokens were stored in a database, a rogue administrator could steal Bob's token and plant it in a forged leak. By storing only C, innocent users cannot be framed."),
        ("Binding Property: ", "Once C is mined on-chain, the issuer cannot alter the token associated with that download.")
    ]
    for label, desc in commit_points:
        p = tf7_2.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    # =========================================================================
    # SLIDE 8: INVISIBLE STEGANOGRAPHY (ZERO-WIDTH UNICODE)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Forensic Watermarking: Zero-Width Unicode Steganography")

    c8_1, tf8_1 = add_card(s8, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="Invisible Code Point Mapping",
                           subtitle="Converting 256 Bits into Non-Printing Characters",
                           border_color=C_CYAN)
    stego_points = [
        ("\\u200B (Zero Width Space): ", "Represents binary bit '0'."),
        ("\\u200C (Zero Width Non-Joiner): ", "Represents binary bit '1'."),
        ("\\uFEFF (Zero Width No-Break Space / BOM): ", "START frame delimiter."),
        ("\\u2060 (Word Joiner): ", "END frame delimiter."),
        ("Frame Payload: ", "1 Start + 256 Bits + 1 End = 258 invisible characters."),
        ("Human Perception: ", "Completely invisible to human readers; renders 0 width in browsers, Word, PDF, and text editors.")
    ]
    for label, desc in stego_points:
        p = tf8_1.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    c8_2, tf8_2 = add_card(s8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
                           title="Multi-Paragraph Redundancy",
                           subtitle="Resilience Against Partial Text Leaks",
                           border_color=C_PRIMARY)
    redundancy_points = [
        ("The Exfiltration Reality: ", "Leakers rarely copy an entire 50-page document. They copy a single juicy paragraph."),
        ("Our Defense Strategy: ", "The engine embeds a full 258-character marker frame at the end of EVERY paragraph."),
        ("Full Leak Recovery: ", "If user copies 100% of the document -> 100% extracted."),
        ("Partial Leak Recovery: ", "If user copies ONLY Paragraph 2 -> The entire 256-bit token is still 100% recovered and matched on-chain!"),
        ("Parser Robustness: ", "Handles truncated frames, malformed bits, and strips invisible characters on demand.")
    ]
    for label, desc in redundancy_points:
        p = tf8_2.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    # =========================================================================
    # SLIDE 9: ATOMIC ENFORCEMENT & DOWNLOAD WORKFLOW
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Atomic Enforcement: The Verified Download Pipeline")

    steps = [
        ("Step 1: Auth Verification", "JWT Bearer token authenticated; resolves user's Ethereum wallet address."),
        ("Step 2: Live On-Chain Check", "Backend queries DataLeakLedger.hasAccess(docId, user). If revoked -> HTTP 403 Forbidden."),
        ("Step 3: Cryptographic Issuance", "Generates 256-bit CSPRNG token T and computes commitment C = Keccak256(T)."),
        ("Step 4: Mined On-Chain Receipt", "Submits recordIssuance(C, docId, version, user). Awaits receipt (status == 1)."),
        ("Step 5: Dynamic Watermarking", "Embeds token T into text/PDF using zero-width characters and streams file to client.")
    ]
    card_h = Inches(0.9)
    top_pos = Inches(1.5)
    for i, (title_st, desc_st) in enumerate(steps):
        c_step, tf_step = add_card(s9, Inches(0.8), top_pos + Inches(i * 1.05), Inches(11.73), card_h,
                                   title=f"{title_st} — {desc_st}",
                                   subtitle="",
                                   border_color=C_PRIMARY if i % 2 == 0 else C_CYAN)

    # =========================================================================
    # SLIDE 10: FORENSIC EXTRACTION & ATTRIBUTION
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Forensic Attribution: Catching the Leaker Step-by-Step")

    c10_1, tf10_1 = add_card(s10, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2),
                             title="1. Extraction Pipeline",
                             subtitle="Reconstructing the 256-Bit Token",
                             border_color=C_CYAN)
    ext_steps = [
        ("Delimiter Detection: ", "Parser scans pasted text for \\uFEFF (START) and \\u2060 (END)."),
        ("Bit Stream Parsing: ", "\\u200B -> '0', \\u200C -> '1'. Validates bit length == 256."),
        ("Hex Reconstruction: ", "Converts 256 bits into 32-byte hex token string T."),
        ("One-Way Hash: ", "Computes C = Keccak256(T)."),
        ("EVM Query: ", "Calls DataLeakLedger.getIssuance(C).")
    ]
    for label, desc in ext_steps:
        p = tf10_1.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    c10_2, tf10_2 = add_card(s10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
                             title="2. Positive Attribution Match",
                             subtitle="Irrefutable Forensic Evidence Returned",
                             border_color=C_EMERALD)
    match_items = [
        ("Recipient Identity: ", "Returns exact wallet address (e.g. Alice: 0x3C44...93BC) and username."),
        ("Document Trace: ", "Identifies document ID and exact version number (e.g. v1)."),
        ("Proof of Time: ", "Returns exact block timestamp when the copy was issued."),
        ("Transaction Audit: ", "Returns mined transaction hash on the Ethereum ledger.")
    ]
    for label, desc in match_items:
        p = tf10_2.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    # =========================================================================
    # SLIDE 11: LIVE DEMONSTRATION WORKFLOW (7 STEPS)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Live Demonstration Walkthrough: The 7-Step Viva Script")

    demo_steps = [
        ("Step 1: Admin Registers Document", "Uploads sample-policy.md; canonical SHA-256 hash mined in Block #1."),
        ("Step 2: Admin Grants Access to Alice", "Calls grantAccess(docId, aliceAddress) on DataLeakLedger contract."),
        ("Step 3: Alice Downloads File", "Backend verifies contract state, commits Keccak-256 hash, and watermarks file."),
        ("Step 4: Alice Simulates Text Leak", "Copies a text snippet containing invisible zero-width characters."),
        ("Step 5: Forensic Match", "Admin pastes leaked text into Forensic Portal -> Matched to Alice's wallet & Block #X!"),
        ("Step 6: Revoke Access & Block", "Admin revokes Alice; WebSockets update Alice's UI to NO ACCESS. Next download blocked!"),
        ("Step 7: Blockchain Explorer", "Inspect blocks, hashes, transaction receipts, and emitted event logs live on-chain.")
    ]
    card_h = Inches(0.68)
    top_pos = Inches(1.5)
    for i, (title_st, desc_st) in enumerate(demo_steps):
        c_step, tf_step = add_card(s11, Inches(0.8), top_pos + Inches(i * 0.76), Inches(11.73), card_h,
                                   title=f"{title_st}: {desc_st}",
                                   subtitle="",
                                   border_color=C_PRIMARY if i % 2 == 0 else C_CYAN)

    # =========================================================================
    # SLIDE 12: LIMITATIONS, VIVA DEFENSE & CONCLUSION
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Viva Defense: Limitations & Enterprise Roadmap")

    c12_1, tf12_1 = add_card(s12, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2),
                             title="What It Proves vs. Limitations",
                             subtitle="Academic Honesty & Boundary Clarity",
                             border_color=C_AMBER)
    limits = [
        ("What It Mathematically Proves: ", "The leaked text contains a cryptographic token whose commitment was mined on-chain when a specific copy was issued to Alice's wallet."),
        ("The Analog Hole: ", "Photographing a computer screen or manually re-typing text on another keyboard bypasses zero-width Unicode characters."),
        ("Text Normalization: ", "Running text through regex filters that explicitly strip non-printable characters removes the watermark."),
        ("Device Ownership: ", "Proves the issuance was signed to Alice's wallet; cannot prove physical theft of Alice's laptop.")
    ]
    for label, desc in limits:
        p = tf12_1.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    c12_2, tf12_2 = add_card(s12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
                             title="Consensus & Enterprise Roadmap",
                             subtitle="Moving Beyond Local Hardhat to QBFT",
                             border_color=C_EMERALD)
    roadmap = [
        ("Hardhat Distinction: ", "Hardhat local node simulates EVM execution and transaction mining for development, but does not run multi-validator Byzantine consensus."),
        ("Consortium Roadmap: ", "Deploy on Hyperledger Besu or GoQuorum running QBFT (Quorum Byzantine Fault Tolerance)."),
        ("Independent Validators: ", "Auditors, Regulators, and Legal Counsel operate distinct validator nodes."),
        ("Conclusion: ", "DataLeak Ledger establishes a practical, highly demonstrable paradigm for tamper-evident document governance and forensic accountability.")
    ]
    for label, desc in roadmap:
        p = tf12_2.add_paragraph()
        p.text = f"• {label}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_LIGHT

    # Save presentation
    output_path = "/Users/johaan/Documents/Projects/blockchain/DataLeak_Ledger_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
